"""Successful local conversions for the formats acceptance covers."""

from __future__ import annotations

import pytest

from helpers import call


@pytest.mark.parametrize(
    "name,needle",
    [
        ("sample.pdf", "Sample Report"),
        ("sample.docx", "Sample Report"),
        ("sample.xlsx", "widget"),
        ("sample.html", "Sample Report"),
    ],
)
def test_convert_formats(plugin, fixtures_dir, name, needle):
    out = call(plugin, "doc_convert", {"path": str(fixtures_dir / name)})
    assert out.get("ok") is True, out
    assert needle in out["markdown"]
    assert out["markitdown_plugins"] is False
    assert out["llm_helpers"] is False


def test_convert_pptx_and_csv(plugin, fixtures_dir):
    pptx = call(plugin, "doc_convert", {"path": str(fixtures_dir / "sample.pptx")})
    assert pptx.get("ok") is True
    assert "Sample Report" in pptx["markdown"]
    csv = call(plugin, "doc_convert", {"path": str(fixtures_dir / "sample.csv")})
    assert csv.get("ok") is True
    assert "gadget" in csv["markdown"]


def test_max_chars_truncates(plugin, fixtures_dir):
    out = call(
        plugin,
        "doc_convert",
        {"path": str(fixtures_dir / "sample.html"), "max_chars": 20},
    )
    assert out.get("ok") is True
    assert out["truncated"] is True
    assert len(out["markdown"]) == 20
    assert out["char_count"] == 20


def _bind_output_chars(plugin, limit: int):
    plugin.settings.bind_config(
        lambda key, default=None: limit if key == "max_output_chars" else default
    )


def test_max_chars_cannot_exceed_config(plugin, fixtures_dir):
    """config 50: omitted uses 50; 10 cuts at 10; 400 is not an error and cuts at 50."""
    _bind_output_chars(plugin, 50)
    try:
        path = str(fixtures_dir / "sample.html")
        omitted = call(plugin, "doc_convert", {"path": path})
        smaller = call(plugin, "doc_convert", {"path": path, "max_chars": 10})
        larger = call(plugin, "doc_convert", {"path": path, "max_chars": 400})
    finally:
        plugin.settings.bind_config(None)
    assert omitted.get("ok") is True, omitted
    assert omitted["truncated"] is True
    assert len(omitted["markdown"]) == 50
    assert omitted["char_count"] == 50
    assert smaller.get("ok") is True, smaller
    assert len(smaller["markdown"]) == 10
    assert smaller["truncated"] is True
    assert larger.get("ok") is True, larger
    assert larger.get("error_type") is None
    assert len(larger["markdown"]) == 50
    assert larger["truncated"] is True
    assert larger["char_count"] == 50


def test_max_chars_zero_and_non_integer(plugin, fixtures_dir):
    path = str(fixtures_dir / "sample.txt")
    zero = call(plugin, "doc_convert", {"path": path, "max_chars": 0})
    negative = call(plugin, "doc_convert", {"path": path, "max_chars": -3})
    text = call(plugin, "doc_convert", {"path": path, "max_chars": "10"})
    assert zero["error_type"] == "bad_argument"
    assert negative["error_type"] == "bad_argument"
    assert text["error_type"] == "bad_argument"


def test_text_json_and_md_convert(plugin, tmp_path):
    text = tmp_path / "notes.txt"
    text.write_text("plain fixture line\n", encoding="utf-8")
    md = tmp_path / "notes.md"
    md.write_text("# Heading\n\nbody line\n", encoding="utf-8")
    js = tmp_path / "notes.json"
    js.write_text('{"widget": "gadget"}\n', encoding="utf-8")
    for path, needle in (
        (text, "plain fixture line"),
        (md, "Heading"),
        (js, "gadget"),
    ):
        out = call(plugin, "doc_convert", {"path": str(path)})
        assert out.get("ok") is True, (path.name, out)
        assert needle in out["markdown"]


_CP932_PHRASE = "①髙﨑～－"


def test_local_html_meta_shift_jis_keeps_cp932_glyphs(plugin, tmp_path):
    html = (
        '<html><head><meta charset="shift_jis"></head><body><p>'
        + _CP932_PHRASE
        + "</p></body></html>"
    )
    path = tmp_path / "page.html"
    path.write_bytes(html.encode("cp932"))
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is True, out
    assert _CP932_PHRASE in out["markdown"]
    assert "郄" not in out["markdown"]


def test_local_html_does_not_rewrite_charset_in_the_body(plugin, tmp_path):
    html = (
        '<html><head><meta charset="shift_jis"></head><body><p>'
        + _CP932_PHRASE
        + " charset=shift_jis</p></body></html>"
    )
    path = tmp_path / "page.html"
    path.write_bytes(html.encode("cp932"))
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is True, out
    assert _CP932_PHRASE in out["markdown"]
    shown = out["markdown"].replace("\\_", "_")
    assert "charset=shift_jis" in shown
    assert "charset=utf-8" not in shown


def test_unknown_html_charset_is_not_a_successful_conversion(plugin, tmp_path):
    path = tmp_path / "page.html"
    path.write_text(
        '<html><head><meta charset="not-a-charset"></head><body><p>Hello</p></body></html>',
        encoding="utf-8",
    )
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is not True
    assert out["error_type"] == "encoding_unconfirmed"
    assert "not-a-charset" in out["error"]
    assert out.get("markdown") in (None, "")


def test_empty_rss_report_does_not_succeed(plugin, tmp_path, monkeypatch):
    import sys

    if sys.platform != "darwin":
        return

    monkeypatch.setattr(plugin.convert.subprocess, "check_output", lambda *_a, **_k: "")
    path = tmp_path / "note.txt"
    path.write_text("hello empty ps", encoding="utf-8")
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is not True
    assert out["error_type"] == "memory_watch_failed"
    assert out.get("markdown") in (None, "")


def test_unreadable_rss_watch_does_not_succeed(plugin, tmp_path, monkeypatch):
    import subprocess
    import sys

    if sys.platform != "darwin":
        return
    def _boom(*_args, **_kwargs):
        raise subprocess.CalledProcessError(1, ["ps"])

    monkeypatch.setattr(plugin.convert.subprocess, "check_output", _boom)
    path = tmp_path / "note.txt"
    path.write_text("hello rss", encoding="utf-8")
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is not True
    assert out["error_type"] == "memory_watch_failed"
    assert out.get("markdown") in (None, "")


def test_local_cp932_txt_is_not_a_successful_conversion(plugin, tmp_path):
    path = tmp_path / "note.txt"
    path.write_bytes(_CP932_PHRASE.encode("cp932"))
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is not True
    assert out["error_type"] == "encoding_unconfirmed"
    assert out.get("markdown") in (None, "")
    assert "嘆" not in out["error"]


def test_lied_central_directory_size_is_refused(plugin, tmp_path, monkeypatch):
    """A local header that is larger than the central directory still counts."""
    import zipfile

    cap = 1000
    plugin.settings.bind_config(
        lambda key, default=None: cap if key == "max_file_bytes" else default
    )
    path = tmp_path / "lied.docx"
    member = "word/document.xml"
    payload = b"A" * 40_000
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(member, payload)
    data = bytearray(path.read_bytes())
    signature = b"PK\x01\x02"
    start = 0
    found = False
    encoded = member.encode("ascii")
    while True:
        index = data.find(signature, start)
        assert index >= 0
        name_len = int.from_bytes(data[index + 28 : index + 30], "little")
        name = bytes(data[index + 46 : index + 46 + name_len])
        if name == encoded:
            data[index + 24 : index + 28] = (20).to_bytes(4, "little")
            found = True
            break
        start = index + 4
    assert found
    path.write_bytes(data)
    _forbid_worker(monkeypatch, plugin)
    try:
        out = call(plugin, "doc_convert", {"path": str(path)})
    finally:
        plugin.settings.bind_config(None)
    assert out.get("ok") is not True
    assert out["error_type"] == "uncompressed_too_large", out
    assert "A" * 20 not in out["error"]


def test_empty_txt_is_empty_output(plugin, tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("", encoding="utf-8")
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out["error_type"] == "empty_output"
    assert "next_step" in out
    assert out["error"] == "Conversion produced no Markdown text."


def _forbid_worker(monkeypatch, plugin):
    def boom(*_args, **_kwargs):
        raise AssertionError("worker must not start")

    monkeypatch.setattr(plugin.convert, "start_conversion_worker", boom)


def _zip_bomb(path, cap: int, member: str) -> None:
    import zipfile

    payload = b"A" * (cap + 1)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(member, payload)


def test_uncompressed_office_over_cap_does_not_start_worker(plugin, tmp_path, monkeypatch):
    cap = 400
    plugin.settings.bind_config(
        lambda key, default=None: cap if key == "max_file_bytes" else default
    )
    path = tmp_path / "bomb.docx"
    _zip_bomb(path, cap, "word/document.xml")
    assert path.stat().st_size <= cap
    _forbid_worker(monkeypatch, plugin)
    try:
        out = call(plugin, "doc_convert", {"path": str(path)})
    finally:
        plugin.settings.bind_config(None)
    assert out["error_type"] == "uncompressed_too_large"
    assert "next_step" in out
    assert "A" * 20 not in out["error"]
    assert out.get("ok") is not True


def test_zip_bodied_xls_is_refused_before_conversion(plugin, tmp_path, monkeypatch):
    path = tmp_path / "bomb.xls"
    _zip_bomb(path, 400, "xl/workbook.xml")
    _forbid_worker(monkeypatch, plugin)
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out["error_type"] == "zip_refused"
    assert out.get("ok") is not True
    assert "Content from the zip file" not in str(out)


def test_non_zip_xls_is_not_uncompressed_too_large(plugin, tmp_path):
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"not-a-zip-xls")
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is not True
    assert out.get("error_type") == "conversion_failed"
    assert "Content from the zip file" not in str(out)


def _nested_docx(path, *, inner_chars: int, wraps: int) -> None:
    """Wrap payload in `wraps` zip layers. wraps=3 is the deep archive case."""
    import io
    import zipfile

    blob = b"x" * inner_chars
    name = "big.txt"
    for _ in range(wraps):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(name, blob)
        blob = buf.getvalue()
        name = "inner.zip"
    path.write_bytes(blob)


def test_deep_nested_docx_is_not_ok(plugin, tmp_path, monkeypatch):
    """315-byte-class deep archive named .docx must not come back ok."""
    path = tmp_path / "deep.docx"
    _nested_docx(path, inner_chars=2_000_000, wraps=3)
    assert path.stat().st_size < 1000
    _forbid_worker(monkeypatch, plugin)
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is not True
    assert out["error_type"] == "zip_refused"
    assert "Content from the zip file" not in str(out)


def test_zip_renamed_to_docx_is_not_ok(plugin, tmp_path, monkeypatch):
    import zipfile

    path = tmp_path / "renamed.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("hello.txt", "not an office package")
    _forbid_worker(monkeypatch, plugin)
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is not True
    assert out["error_type"] == "zip_refused"
    assert "Content from the zip file" not in str(out)


def test_zip_extension_is_unsupported_before_conversion(plugin, tmp_path, monkeypatch):
    import zipfile

    path = tmp_path / "plain.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("hello.txt", "zip")
    _forbid_worker(monkeypatch, plugin)
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is not True
    assert out["error_type"] == "unsupported_type"


def test_zip_renamed_to_html_is_refused(plugin, tmp_path, monkeypatch):
    import zipfile

    path = tmp_path / "page.html"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("hello.txt", "not html")
    _forbid_worker(monkeypatch, plugin)
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is not True
    assert out["error_type"] == "zip_refused"


def test_invalid_zip_docx_is_conversion_failed_without_worker(plugin, tmp_path, monkeypatch):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"this is not a zip")
    _forbid_worker(monkeypatch, plugin)
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out["error_type"] == "conversion_failed"
    assert "next_step" in out
    assert "this is not a zip" not in out["error"]


def test_dispatch_source_never_selects_an_archive_walk():
    from helpers import PLUGIN_DIR

    for name in ("convert.py", "convert_worker.py"):
        blob = (PLUGIN_DIR / name).read_text(encoding="utf-8")
        assert "convert_local(" not in blob
        assert ".convert_local" not in blob
        assert "convert_stream" not in blob
        assert "ZipConverter" not in blob


def test_success_json_puts_cut_fields_before_markdown(plugin, tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("A" * 4000, encoding="utf-8")
    raw = plugin.tools.HANDLERS["doc_convert"]({"path": str(path), "max_chars": 20})
    preview = raw[:1500]
    assert '"truncated": true' in preview
    assert '"char_count"' in preview
    assert '"error"' in preview
    mark_at = raw.find('"markdown"')
    prefix = raw[:mark_at]
    assert '"truncated"' in prefix
    assert '"char_count"' in prefix
    assert '"error"' in prefix
    assert preview.index('"truncated"') < preview.index('"markdown"')


def test_chart_pptx_with_embedded_xlsx_converts(plugin, tmp_path):
    import zipfile

    pptx = pytest.importorskip("pptx")
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    path = tmp_path / "chart.pptx"
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    chart_data = CategoryChartData()
    chart_data.categories = ["Apple", "Banana"]
    chart_data.add_series("Sales", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1),
        Inches(6),
        Inches(4),
        chart_data,
    )
    presentation.save(str(path))
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    assert "ppt/presentation.xml" in names
    assert any(name.endswith(".xlsx") for name in names)
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is True, out
    assert "Content from the zip file" not in out["markdown"]
    assert "Apple" in out["markdown"]
    assert "Sales" in out["markdown"]


def test_real_pdf_and_html_still_convert(plugin, fixtures_dir):
    pdf = call(plugin, "doc_convert", {"path": str(fixtures_dir / "sample.pdf")})
    html = call(plugin, "doc_convert", {"path": str(fixtures_dir / "sample.html")})
    assert pdf.get("ok") is True, pdf
    assert "Sample Report" in pdf["markdown"]
    assert html.get("ok") is True, html
    assert "Sample Report" in html["markdown"]
    assert pdf["error"] is None
    assert html["truncated"] is False


def test_one_second_limit_still_converts_a_small_file(plugin, tmp_path):
    plugin.settings.bind_config(
        lambda key, default=None: 1 if key == "convert_timeout_seconds" else default
    )
    path = tmp_path / "note.txt"
    path.write_text("hello small", encoding="utf-8")
    try:
        out = call(plugin, "doc_convert", {"path": str(path)})
    finally:
        plugin.settings.bind_config(None)
    assert out.get("ok") is True, out
    assert "hello small" in out["markdown"]


def test_five_megabyte_html_stops_at_the_time_limit(plugin, tmp_path):
    path = tmp_path / "heavy.html"
    chunk = b"<b>a</b>" * 10000
    target = 5 * 1024 * 1024
    with path.open("wb") as handle:
        handle.write(b"<html><body>")
        while handle.tell() < target:
            handle.write(chunk)
        handle.write(b"</body></html>")
    assert path.stat().st_size >= target
    out = call(plugin, "doc_convert", {"path": str(path)})
    limit = plugin.settings.DEFAULT_CONVERT_TIMEOUT
    assert out.get("ok") is not True
    assert out["error_type"] == "conversion_timeout", out
    assert f"time limit of {limit} seconds" in out["error"]
    assert "stopped" in out["error"]


def _assert_stopped_for_memory(out):
    assert out.get("ok") is not True, out
    assert out.get("markdown") in (None, "")
    assert out["error_type"] == "conversion_memory", out
    assert "922746880" in out["error"]
    assert "memory limit" in out["error"]


def test_pk_comma_csv_converts(plugin, tmp_path):
    path = tmp_path / "people.csv"
    path.write_text("PK,Name,Email\n1,Ada,ada@example.com\n", encoding="utf-8")
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is True, out
    assert "Ada" in out["markdown"]


def test_pki_markdown_converts(plugin, tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("PKI rollout notes\n", encoding="utf-8")
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is True, out
    assert "rollout" in out["markdown"]


def test_child_does_not_receive_hermes_or_key_env(plugin, tmp_path, monkeypatch):
    import subprocess

    monkeypatch.setenv("OPENAI_API_KEY", "canary-openai")
    monkeypatch.setenv("HERMES_API_KEY", "canary-hermes")
    monkeypatch.setenv("AWS_SECRET_ACCESS_KEY", "canary-aws")
    seen = {}
    real_popen = subprocess.Popen

    def wrapped(cmd, *args, **kwargs):
        if isinstance(cmd, (list, tuple)) and any("convert_worker.py" in str(part) for part in cmd):
            seen["env"] = dict(kwargs.get("env") or {})
        return real_popen(cmd, *args, **kwargs)

    monkeypatch.setattr(plugin.convert.subprocess, "Popen", wrapped)
    path = tmp_path / "note.txt"
    path.write_text("hello env\n", encoding="utf-8")
    out = call(plugin, "doc_convert", {"path": str(path)})
    assert out.get("ok") is True, out
    env = seen["env"]
    assert "OPENAI_API_KEY" not in env
    assert "HERMES_API_KEY" not in env
    assert "HERMES_HOME" not in env
    assert "AWS_SECRET_ACCESS_KEY" not in env
    assert "canary-openai" not in out["markdown"]
    assert env.get("PYTHONUNBUFFERED") == "1"
    assert "PATH" in env


def test_linux_address_limit_is_rlimit_as(plugin, monkeypatch):
    import resource

    calls = {}

    def fake_setrlimit(which, limits):
        calls["which"] = which
        calls["limits"] = limits

    monkeypatch.setattr(resource, "setrlimit", fake_setrlimit)
    monkeypatch.setattr(plugin.convert.sys, "platform", "linux")
    plugin.convert.apply_linux_address_limit()
    assert calls["which"] == resource.RLIMIT_AS
    limit = plugin.convert.CONVERT_MEMORY_LIMIT_BYTES
    assert calls["limits"] == (limit, limit)
    assert limit == 922746880
    calls.clear()
    monkeypatch.setattr(plugin.convert.sys, "platform", "darwin")
    plugin.convert.apply_linux_address_limit()
    assert calls == {}


def _sparse_xlsx(path):
    import io
    import zipfile

    sheet = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        "<sheetData><row r=\"1\"><c r=\"A1\"><v>1</v></c></row>"
        "<row r=\"1048576\"><c r=\"XFD1048576\"><v>2</v></c></row>"
        "</sheetData></worksheet>"
    )
    parts = {
        "[Content_Types].xml": (
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/>'
            '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
            '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
            "</Types>"
        ),
        "_rels/.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
            "</Relationships>"
        ),
        "xl/workbook.xml": (
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="S" sheetId="1" r:id="rId1"/></sheets></workbook>'
        ),
        "xl/_rels/workbook.xml.rels": (
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
            "</Relationships>"
        ),
        "xl/worksheets/sheet1.xml": sheet,
    }
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, body in parts.items():
            archive.writestr(name, body)
    path.write_bytes(buf.getvalue())


def test_sparse_xlsx_stops_at_the_memory_limit(plugin, tmp_path):
    path = tmp_path / "sparse.xlsx"
    _sparse_xlsx(path)
    size = path.stat().st_size
    assert 1200 <= size <= 1800, size
    out = call(plugin, "doc_convert", {"path": str(path)})
    _assert_stopped_for_memory(out)


def _patch_zip_sizes(data, name, declared, crc):
    def _patch(signature, name_at, crc_at, size_at):
        start = 0
        found = False
        while True:
            index = data.find(signature, start)
            if index < 0:
                break
            name_len = int.from_bytes(data[index + name_at - 4 : index + name_at - 2], "little")
            # local name length is at +26; central name length is at +28.
            if signature == b"PK\x03\x04":
                name_len = int.from_bytes(data[index + 26 : index + 28], "little")
                got = bytes(data[index + 30 : index + 30 + name_len])
            else:
                name_len = int.from_bytes(data[index + 28 : index + 30], "little")
                got = bytes(data[index + 46 : index + 46 + name_len])
            if got == name:
                data[index + crc_at : index + crc_at + 4] = crc.to_bytes(4, "little")
                data[index + size_at : index + size_at + 4] = declared.to_bytes(4, "little")
                found = True
                break
            start = index + 4
        assert found

    _patch(b"PK\x03\x04", 30, 14, 22)
    _patch(b"PK\x01\x02", 46, 16, 24)


def _forged_docx(path):
    """Both headers understate a longer deflate stream of paragraph XML."""
    import io
    import zipfile
    import zlib

    para = b"<w:p><w:r><w:t>A</w:t></w:r></w:p>"
    head = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
    )
    tail = b"</w:body></w:document>"
    count = (18 * 1024 * 1024 - len(head) - len(tail)) // len(para)
    declared_xml = head + (para * count) + tail
    full = declared_xml + (para * 40000)
    assert len(full) > len(declared_xml)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            "</Types>",
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>",
        )
        archive.writestr("word/document.xml", full)
    data = bytearray(buf.getvalue())
    declared = len(declared_xml)
    crc = zlib.crc32(declared_xml) & 0xFFFFFFFF
    _patch_zip_sizes(data, b"word/document.xml", declared, crc)
    path.write_bytes(data)
    with zipfile.ZipFile(path) as archive:
        info = archive.getinfo("word/document.xml")
        assert info.file_size == declared
        assert info.file_size < len(full)
        got = archive.read("word/document.xml")
        assert got == declared_xml


def test_forged_docx_sizes_stop_at_the_memory_limit(plugin, tmp_path):
    path = tmp_path / "forged.docx"
    _forged_docx(path)
    out = call(plugin, "doc_convert", {"path": str(path)})
    _assert_stopped_for_memory(out)


def _flate_pdf(path):
    """A small PDF whose FlateDecode stream inflates to about 3 GiB of text.

    FlateDecode is the zlib wrapper (RFC 1950), not raw deflate. The page
    uses a real font resource so the decompressed text is retained.
    """
    import zlib

    compressor = zlib.compressobj(9)
    chunks = [compressor.compress(b"BT /F1 24 Tf 72 720 Td (")]
    block = b"A" * (1024 * 1024)
    for _ in range(3072):
        chunks.append(compressor.compress(block))
    chunks.append(compressor.compress(b") Tj ET"))
    chunks.append(compressor.flush())
    stream = b"".join(chunks)
    obj1 = b"1 0 obj\n<< /Type /Catalog /Pages 2 0 R >>\nendobj\n"
    obj2 = b"2 0 obj\n<< /Type /Pages /Kids [3 0 R] /Count 1 >>\nendobj\n"
    obj3 = (
        b"3 0 obj\n<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\nendobj\n"
    )
    obj4 = (
        b"4 0 obj\n<< /Length "
        + str(len(stream)).encode()
        + b" /Filter /FlateDecode >>\nstream\n"
        + stream
        + b"\nendstream\nendobj\n"
    )
    obj5 = b"5 0 obj\n<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\nendobj\n"
    blob = bytearray(b"%PDF-1.4\n")
    offsets = []
    for obj in (obj1, obj2, obj3, obj4, obj5):
        offsets.append(len(blob))
        blob += obj
    xref_at = len(blob)
    xref = bytearray(b"xref\n0 6\n0000000000 65535 f \n")
    for offset in offsets:
        xref += f"{offset:010d} 00000 n \n".encode()
    trailer = (
        b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n"
        + str(xref_at).encode()
        + b"\n%%EOF\n"
    )
    path.write_bytes(bytes(blob) + bytes(xref) + trailer)


def test_flate_pdf_stops_at_the_memory_limit(plugin, tmp_path):
    path = tmp_path / "bomb.pdf"
    _flate_pdf(path)
    assert path.stat().st_size < 5 * 1024 * 1024
    out = call(plugin, "doc_convert", {"path": str(path)})
    _assert_stopped_for_memory(out)
